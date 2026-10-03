"""Excel/CSV parsing and validation for student imports."""

import csv
import io
import re
from datetime import date, datetime
from typing import IO, Any


HEADER_ALIASES = {
    "full_name": {"full name", "student name", "name"},
    "gender": {"gender"},
    "date_of_birth": {"date of birth", "date of birth (yyyy-mm-dd)", "dob", "dob (yyyy-mm-dd)"},
    "parent_name": {"parent/guardian name", "parent name", "guardian name"},
    "parent_contact": {"parent contact number", "parent contact", "contact", "phone"},
    "parent_name_2": {"secondary contact name (optional)", "secondary contact name", "second parent name"},
    "parent_contact_2": {"whatsapp number", "whatsapp number (optional)", "secondary contact number (optional)", "secondary contact (optional)", "secondary contact", "second parent contact"},
    "class_name": {"class", "assign to class", "class name"},
    "joined_date": {"joined date", "joined date (yyyy-mm-dd)", "date joined"},
    # Legacy template columns remain supported.
    "grade": {"grade"},
    "medium": {"medium", "medium (sinhala/tamil)"},
}

TEMPLATE_HEADERS = [
    "Full Name",
    "Gender",
    "Date of Birth (YYYY-MM-DD)",
    "Parent/Guardian Name",
    "Parent Contact Number",
    "Secondary Contact Name (Optional)",
    "WhatsApp Number (Optional)",
    "Class",
    "Joined Date (YYYY-MM-DD)",
]


def create_student_import_template(classes: list[str]) -> io.BytesIO:
    """Create a styled workbook with sample rows and current class dropdowns."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Student Import"
    options = workbook.create_sheet("Options")

    header_fill = PatternFill("solid", fgColor="059669")
    required_fill = PatternFill("solid", fgColor="F8FAFC")
    optional_fill = PatternFill("solid", fgColor="FEF3C7")
    error_fill = PatternFill("solid", fgColor="FEE2E2")

    for column, header in enumerate(TEMPLATE_HEADERS, start=1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [28, 13, 24, 26, 22, 31, 33, 32, 24]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    worksheet.row_dimensions[1].height = 34
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:I500"

    for row in range(2, 501):
        for column in range(1, 10):
            worksheet.cell(row=row, column=column).fill = optional_fill if column in (6, 7) else required_fill
        worksheet.cell(row=row, column=5).number_format = "@"
        worksheet.cell(row=row, column=7).number_format = "@"
        worksheet.cell(row=row, column=3).number_format = "yyyy-mm-dd"
        worksheet.cell(row=row, column=9).number_format = "yyyy-mm-dd"

    # Use available classes so the examples match this school's dropdown values.
    def example_class(gender: str) -> str:
        return next(
            (name for name in classes if name.endswith((f" {gender}", " Mixed"))),
            classes[0] if classes else "",
        )

    examples = [
        ["Sample Student Ahmed", "Male", date(2015, 5, 10), "Sample Guardian Ali",
         "0771234567", "Sample Guardian Amina", "0712345678", example_class("Male"), date(2026, 1, 11)],
        ["Sample Student Fathima", "Female", date(2016, 8, 20), "Sample Guardian Mohamed",
         "0751234567", None, None, example_class("Female"), date(2026, 1, 11)],
    ]
    for row_number, values in enumerate(examples, start=2):
        for column, value in enumerate(values, start=1):
            worksheet.cell(row=row_number, column=column, value=value)
        worksheet.cell(row=row_number, column=1).comment = Comment(
            "Example only. Replace or delete both sample rows (2 and 3) before uploading.",
            "Student Import",
        )

    instructions = workbook.create_sheet("Instructions")
    for instruction in [
        "How to fill the Student Import sheet",
        "Rows 2 and 3 are fictional examples. Replace or delete BOTH rows before uploading; otherwise they will be imported.",
        "Keep row 1 (the column headers) unchanged. Enter one student per row, starting at row 2.",
        "Enter dates as YYYY-MM-DD, for example 2015-05-10. Both Date of Birth and Joined Date are required.",
        "Choose Male or Female from the Gender dropdown and the student's current class from the Class dropdown.",
        "Contact numbers must contain 10 digits, including the leading zero, for example 0771234567.",
        "Secondary Contact Name and WhatsApp Number are optional; row 3 shows these left blank.",
        "All other columns are required. If there are no Class options, create current-year classes and download a new template.",
    ]:
        instructions.append([instruction])
    instructions.column_dimensions["A"].width = 110
    for row in instructions:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        instructions.row_dimensions[row[0].row].height = 32
    instructions["A1"].font = Font(bold=True, color="059669", size=14)

    options.append(["Gender", "Class"])
    for index, class_name in enumerate(classes, start=2):
        options.cell(row=index, column=2, value=class_name)
    options.cell(row=2, column=1, value="Male")
    options.cell(row=3, column=1, value="Female")

    class_end = max(2, len(classes) + 1)
    workbook.defined_names.add(DefinedName("GenderOptions", attr_text="'Options'!$A$2:$A$3"))
    workbook.defined_names.add(DefinedName("ClassOptions", attr_text=f"'Options'!$B$2:$B${class_end}"))
    gender_validation = DataValidation(type="list", formula1="GenderOptions", allow_blank=False)
    class_validation = DataValidation(type="list", formula1="ClassOptions", allow_blank=False)
    gender_validation.error = "Choose Male or Female from the list."
    gender_validation.errorTitle = "Invalid gender"
    class_validation.error = "Choose a current class from the list."
    class_validation.errorTitle = "Invalid class"
    for validation, cells in ((gender_validation, "B2:B500"), (class_validation, "H2:H500")):
        validation.errorStyle = "stop"
        validation.showErrorMessage = True
        validation.showInputMessage = True
        worksheet.add_data_validation(validation)
        validation.add(cells)

    # Clearly flag incomplete required fields while users prepare the sheet.
    worksheet.conditional_formatting.add(
        "A2:I500",
        FormulaRule(formula=['AND(COUNTA($A2:$I2)>0,OR($A2="",$B2="",$C2="",$D2="",$E2="",$H2="",$I2=""))'], fill=error_fill),
    )

    options.sheet_state = "hidden"
    workbook.calculation.fullCalcOnLoad = True
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def parse_and_validate_excel(file_obj: IO, filename: str) -> dict:
    """Parse a workbook/CSV into validated, canonical student rows."""
    if filename.lower().endswith(".csv"):
        rows, errors = _parse_csv(file_obj)
    else:
        rows, errors = _parse_xlsx(file_obj)

    return {
        "rows": rows,
        "errors": errors,
        "valid_count": sum(1 for row in rows if not row["errors"]),
    }


def _normalise_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _header_indexes(headers: list[Any]) -> dict[str, int]:
    indexes: dict[str, int] = {}
    for index, header in enumerate(headers):
        normalised = _normalise_header(header)
        for field, aliases in HEADER_ALIASES.items():
            if normalised in aliases and field not in indexes:
                indexes[field] = index
                break
    return indexes


def _is_empty_row(values: list[Any] | tuple[Any, ...]) -> bool:
    return not any(str(value or "").strip() for value in values)


def _parse_csv(file_obj: IO):
    content = file_obj.read().decode("utf-8-sig")
    all_rows = list(csv.reader(io.StringIO(content)))
    return _parse_rows(all_rows)


def _parse_xlsx(file_obj: IO):
    from openpyxl import load_workbook

    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    worksheet = workbook["Student Import"] if "Student Import" in workbook.sheetnames else workbook.active
    all_rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    return _parse_rows(all_rows)


def _parse_rows(all_rows: list[list[Any]] | list[tuple[Any, ...]]):
    if not all_rows:
        return [], [{"row": 0, "field": "file", "message": "Empty file"}]

    headers = list(all_rows[0])
    indexes = _header_indexes(headers)
    required_headers = {"full_name", "gender", "date_of_birth", "parent_name", "parent_contact"}
    missing = sorted(required_headers - indexes.keys())
    if missing:
        labels = ", ".join(field.replace("_", " ").title() for field in missing)
        return [], [{"row": 1, "field": "headers", "message": f"Missing required columns: {labels}"}]

    has_class = "class_name" in indexes
    has_legacy_class = "grade" in indexes and "medium" in indexes
    if not has_class and not has_legacy_class:
        return [], [{"row": 1, "field": "headers", "message": "Missing required Class column"}]

    rows = []
    errors = []
    for row_number, values in enumerate(all_rows[1:], start=2):
        values = list(values)
        if _is_empty_row(values):
            continue

        def value(field: str):
            index = indexes.get(field)
            return values[index] if index is not None and index < len(values) else None

        raw = {field: value(field) for field in HEADER_ALIASES}
        raw["joined_date_header_present"] = "joined_date" in indexes
        parsed, row_errors = _validate_row(raw, row_number)
        rows.append(parsed)
        errors.extend(row_errors)

    if not rows:
        errors.append({"row": 2, "field": "file", "message": "No student rows found"})
    return rows, errors


def _text(value: Any) -> str:
    return str(value or "").strip()


def _parse_date(value: Any, field: str, row_number: int, required: bool = True):
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None

    text = _text(value)
    if not text:
        if required:
            return None, {"row": row_number, "field": field, "message": f"{field.replace('_', ' ').title()} is required"}
        return None, None

    try:
        return datetime.strptime(text.split(" ")[0], "%Y-%m-%d").date(), None
    except ValueError:
        return None, {"row": row_number, "field": field, "message": f"Invalid date '{text}'. Use YYYY-MM-DD"}


def _normalise_contact(value: Any) -> str:
    text = _text(value)
    if text.endswith(".0"):
        text = text[:-2]
    digits = re.sub(r"\D", "", text)
    # Excel often removes the leading zero when a phone number is entered as a number.
    if len(digits) == 9 and digits.startswith("7"):
        digits = f"0{digits}"
    return digits


def _validate_row(raw: dict, row_number: int):
    row_errors = []
    full_name = _text(raw.get("full_name"))
    gender = _text(raw.get("gender")).title()
    parent_name = _text(raw.get("parent_name"))
    parent_name_2 = _text(raw.get("parent_name_2"))
    parent_contact = _normalise_contact(raw.get("parent_contact"))
    parent_contact_2 = _normalise_contact(raw.get("parent_contact_2"))
    class_name = _text(raw.get("class_name"))
    medium = _text(raw.get("medium")).title()
    grade_text = _text(raw.get("grade")).lower().replace("grade", "").strip()

    if not full_name:
        row_errors.append({"row": row_number, "field": "full_name", "message": "Full name is required"})
    if gender not in ("Male", "Female"):
        row_errors.append({"row": row_number, "field": "gender", "message": "Gender must be Male or Female"})
    if not parent_name:
        row_errors.append({"row": row_number, "field": "parent_name", "message": "Parent/guardian name is required"})
    if len(parent_contact) != 10:
        row_errors.append({"row": row_number, "field": "parent_contact", "message": "Parent contact must contain 10 digits"})
    if parent_contact_2 and len(parent_contact_2) != 10:
        row_errors.append({"row": row_number, "field": "parent_contact_2", "message": "WhatsApp number must contain 10 digits"})

    date_of_birth, date_error = _parse_date(raw.get("date_of_birth"), "date_of_birth", row_number)
    if date_error:
        row_errors.append(date_error)

    joined_required = bool(raw.get("joined_date_header_present"))
    joined_date, joined_error = _parse_date(raw.get("joined_date"), "joined_date", row_number, joined_required)
    if joined_error:
        row_errors.append(joined_error)
    if joined_date is None and not joined_required:
        joined_date = date.today()

    grade = None
    if class_name:
        # Grade and medium will be derived from the selected class by the import route.
        pass
    else:
        if medium not in ("Sinhala", "Tamil"):
            row_errors.append({"row": row_number, "field": "medium", "message": "Medium must be Sinhala or Tamil"})
        if not grade_text.isdigit() or not 1 <= int(grade_text) <= 11:
            row_errors.append({"row": row_number, "field": "grade", "message": "Grade must be from 1 to 11"})
        else:
            grade = int(grade_text)

    parsed = {
        "row_num": row_number,
        "full_name": full_name,
        "gender": gender if gender in ("Male", "Female") else None,
        "date_of_birth": date_of_birth,
        "parent_name": parent_name,
        "parent_name_2": parent_name_2 or None,
        "parent_contact": parent_contact,
        "parent_contact_2": parent_contact_2 or None,
        "class_name": class_name or None,
        "joined_date": joined_date,
        "medium": medium if medium in ("Sinhala", "Tamil") else None,
        "grade": grade,
        "errors": [error["message"] for error in row_errors],
    }
    return parsed, row_errors
